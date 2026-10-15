"""A slide with a chart is read, not an ingest error.

A progress-meeting deck in the BFI folder failed with "shape does not contain
a table": a chart is a graphic frame, whose `table` raises instead of being
absent, and the parser tested for the attribute.
"""

from __future__ import annotations

import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from aifactory_rag.config import RagConfig
from aifactory_rag.ingest.parsers import parse_file, parser_version
from aifactory_rag.ingest.pipeline import _ingest_config_matches


def picture(seed: int) -> io.BytesIO:
    """A PNG large enough not to be taken for an icon (random pixels do not compress)."""
    from PIL import Image

    image = Image.frombytes("RGB", (64, 64), os.urandom(64 * 64 * 3) if seed else bytes(64 * 64 * 3))
    data = io.BytesIO()
    image.save(data, format="PNG")
    data.seek(0)
    return data


class PptxParseTests(unittest.TestCase):
    def test_a_chart_is_skipped_and_text_and_tables_are_read(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Project progress"
        data = CategoryChartData()
        data.categories = ["Sep", "Oct"]
        data.add_series("Done", (3, 5))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data)
        table = slide.shapes.add_table(2, 2, Inches(5), Inches(1), Inches(4), Inches(1)).table
        for row, cells in enumerate([("Item", "State"), ("A429 Lite", "done")]):
            for column, text in enumerate(cells):
                table.cell(row, column).text = text

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "progress.pptx"
            presentation.save(path)
            text = parse_file(path)

        self.assertEqual(text, "[slide 1]\nProject progress\nItem | State\nA429 Lite | done")


class PptxPictureTests(unittest.TestCase):
    """The text in a slide's pictures is read, as it is in PDF and DOCX."""

    def test_the_pictures_of_a_slide_are_read_and_a_repeated_one_once(self) -> None:
        presentation = Presentation()
        logo, diagram, screenshot = picture(1).getvalue(), picture(2).getvalue(), picture(3).getvalue()
        first = presentation.slides.add_slide(presentation.slide_layouts[5])
        first.shapes.title.text = "Blocking points"
        first.shapes.add_picture(io.BytesIO(logo), Inches(0), Inches(0))
        group = first.shapes.add_group_shape()
        group.shapes.add_picture(io.BytesIO(diagram), Inches(1), Inches(1))
        second = presentation.slides.add_slide(presentation.slide_layouts[6])
        second.shapes.add_picture(io.BytesIO(logo), Inches(0), Inches(0))
        second.shapes.add_picture(io.BytesIO(screenshot), Inches(2), Inches(2))
        labels = {logo: "BYTECRAFT", diagram: "AXI to GE5 command queue", screenshot: "Vivado timing summary"}

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "blocking-points.pptx"
            presentation.save(path)
            with (
                patch("aifactory_rag.ingest.parsers._ocr_engine", return_value=object()),
                patch("aifactory_rag.ingest.parsers.ocr_text", side_effect=lambda data: labels[data]) as ocr,
            ):
                text = parse_file(path)

        self.assertEqual(text, (
            "[slide 1]\nBlocking points\n(image) BYTECRAFT\nAXI to GE5 command queue\n\n"
            "[slide 2]\n(image) Vivado timing summary"
        ))
        self.assertEqual(ocr.call_count, 3, "the logo on the second slide is not read again")

    def test_without_ocr_a_deck_reads_as_before(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Blocking points"
        slide.shapes.add_picture(picture(1), Inches(0), Inches(0))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "deck.pptx"
            presentation.save(path)
            with patch("aifactory_rag.ingest.parsers._ocr_engine", return_value=None):
                self.assertEqual(parse_file(path), "[slide 1]\nBlocking points")

    def test_a_deck_read_before_its_pictures_were_is_read_again(self) -> None:
        config = RagConfig.model_validate({"embedding": {"provider": "ollama", "model": "m", "dimensions": 3}})
        before = {"metadata": {
            "chunker": "size-v1", "chunkSize": config.ingest.chunk_size, "chunkOverlap": config.ingest.chunk_overlap,
            "embeddingProvider": "ollama", "embeddingModel": "m", "embeddingDimensions": 3,
        }}
        now = {"metadata": {**before["metadata"], "parserVersion": 2}}

        self.assertEqual(parser_version("SOTA/deck.PPTX"), 2)
        self.assertFalse(_ingest_config_matches(before, config, "SOTA/deck.pptx"))
        self.assertTrue(_ingest_config_matches(now, config, "SOTA/deck.pptx"))
        self.assertTrue(_ingest_config_matches(before, config, "SOTA/manual.pdf"), "other types are not read again")


if __name__ == "__main__":
    unittest.main()
