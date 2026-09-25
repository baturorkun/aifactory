"""Spreadsheets are read, not skipped.

A hardware/software interface document (the BFI HSI) arrived as an .xlsx with
one sheet per interface: memory map, UART, ARINC429 register tables. The
ingest had no reader for it and the include list had no pattern, so the corpus
held the drivers written from it but not the document itself, and the RAG
answered that the specification did not exist.
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from aifactory_rag.config import RagSourceConfig
from aifactory_rag.ingest.parsers import SUPPORTED_EXTENSIONS, parse_file
from aifactory_rag.ingest.pipeline import _content_type_metadata as classify


def _hsi_like(path: Path) -> None:
    workbook = Workbook()
    memory = workbook.active
    memory.title = "Memory Map"
    memory.append(["Base Name", "Base Address", "Definition"])
    memory.append(["Fabric Interface Controller Base Address (FIC Base)", "0x3000_0000", None])
    memory.append([])  # an empty row is dropped
    memory.append(["FIC_BASE + 0004_0000", "ARINC429", "Arinc429 Interface (A429 Lite IP)"])

    arinc = workbook.create_sheet("ARINC429")
    arinc.append(["Name", "Bit No", "Access Type", "Initial Value", "Description"])
    arinc.append(["RXFIFOEmpty", 1, "R", 1.0, "RX FIFO empty flag bit.\n0:Not Empty,  1: Empty"])
    arinc.append(["Reserved", "7:4", "R", "0xF", None, None])  # trailing blanks trimmed

    workbook.create_sheet("GPIO")  # a sheet with nothing in it leaves no section
    workbook.save(path)


class XlsxParseTests(unittest.TestCase):
    def test_every_sheet_is_read_with_its_name_and_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "BFI_HSI_v0.2.xlsx"
            _hsi_like(path)
            text = parse_file(path)

        self.assertIn("[sheet Memory Map]", text)
        self.assertIn("FIC_BASE + 0004_0000 | ARINC429 | Arinc429 Interface (A429 Lite IP)", text)
        self.assertIn("[sheet ARINC429]", text)
        # Whole floats lose the .0 and a cell's line break stays on its row.
        self.assertIn("RXFIFOEmpty | 1 | R | 1 | RX FIFO empty flag bit. / 0:Not Empty,  1: Empty", text)
        self.assertIn("Reserved | 7:4 | R | 0xF\n", text + "\n")
        self.assertNotIn("[sheet GPIO]", text)
        self.assertNotIn("\n\n\n", text)

    def test_spreadsheets_are_supported_documentation_and_included_by_default(self) -> None:
        self.assertIn(".xlsx", SUPPORTED_EXTENSIONS)
        self.assertIn(".xlsm", SUPPORTED_EXTENSIONS)
        self.assertEqual(classify("1. RECEIVED/2026.08.26 HSI (v0.2)/BFI_HSI_v0.2.xlsx"),
                         {"contentType": "documentation"})
        include = RagSourceConfig.model_validate({"id": "s", "rootPath": "/x"}).include
        self.assertIn("**/*.xlsx", include)


if __name__ == "__main__":
    unittest.main()
